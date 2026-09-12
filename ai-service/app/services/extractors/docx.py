# ai-service/app/services/extractors/docx.py
from __future__ import annotations

import logging
from pathlib import Path

from docx import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from app.schemas.document import (
    BlockType,
    Document,
    DocumentMetadata,
    DocumentPage,
    DocumentType,
    ImageBlock,
    TableBlock,
    TextBlock,
)
from app.services.extractors.base import BaseExtractor

logger = logging.getLogger(__name__)


class DocxExtractor(BaseExtractor):
    SUPPORTED_EXTENSIONS = {".docx"}

    def extract(
        self,
        file_path: str | Path,
    ) -> Document:

        path = self.validate_file(file_path)
        start = self.log_start(path)

        try:
            doc = DocxDocument(path)

            document = Document(
                metadata=self._extract_metadata(
                    doc,
                    path,
                )
            )

            # python-docx does not provide reliable physical
            # page boundaries. Until layout-aware extraction is
            # added, DOCX content is represented as page 1.
            page = DocumentPage(page_number=1)

            page.blocks.extend(
                self._extract_paragraphs(doc)
            )

            page.blocks.extend(
                self._extract_tables(doc)
            )

            page.blocks.extend(
                self._extract_images(doc)
            )

            document.pages.append(page)

            self.log_success(path, start)
            return document

        except Exception as error:
            self.log_failure(path, error)
            raise

    def _extract_metadata(
        self,
        doc: DocxDocument,
        path: Path,
    ) -> DocumentMetadata:

        props = doc.core_properties

        return DocumentMetadata(
            filename=path.name,
            document_type=DocumentType.DOCX,
            title=props.title or None,
            author=props.author or None,
            metadata={
                "subject": props.subject or None,
                "creator": props.author or None,
                "keywords": props.keywords or None,
                "comments": props.comments or None,
            },
        )

    def _extract_paragraphs(
        self,
        doc: DocxDocument,
    ) -> list[TextBlock]:

        blocks: list[TextBlock] = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()

            if not text:
                continue

            style_name = paragraph.style.name
            style = style_name.lower()

            block_type = (
                BlockType.HEADING
                if "heading" in style
                else BlockType.TEXT
            )

            blocks.append(
                TextBlock(
                    block_id=self.generate_block_id("text"),
                    block_type=block_type,
                    page_number=1,
                    text=text,
                    metadata={
                        "style": style_name,
                    },
                )
            )

        return blocks

    def _extract_tables(
        self,
        doc: DocxDocument,
    ) -> list[TableBlock]:

        blocks: list[TableBlock] = []

        for table in doc.tables:
            rows = [
                [
                    cell.text.strip()
                    for cell in row.cells
                ]
                for row in table.rows
            ]

            if not rows:
                continue

            blocks.append(
                TableBlock(
                    block_id=self.generate_block_id("table"),
                    block_type=BlockType.TABLE,
                    page_number=1,
                    headers=rows[0],
                    rows=rows[1:],
                    metadata={
                        "parser": "python-docx",
                    },
                )
            )

        return blocks

    def _extract_images(
        self,
        doc: DocxDocument,
    ) -> list[ImageBlock]:

        blocks: list[ImageBlock] = []

        for relationship in doc.part.rels.values():

            if relationship.reltype != RT.IMAGE:
                continue

            blocks.append(
                ImageBlock(
                    block_id=self.generate_block_id("image"),
                    block_type=BlockType.IMAGE,
                    page_number=1,
                    metadata={
                        "target": relationship.target_ref,
                    },
                )
            )

        return blocks


docx_extractor = DocxExtractor()